import json
import math
import zipfile
from pathlib import Path

import openpyxl


root = Path(__file__).resolve().parent
results = []
for path in sorted(root.glob("*.xlsx")):
    try:
        zipfile.ZipFile(path).testzip()
    except zipfile.BadZipFile:
        # 非工作簿文件（例如平台返回的 HTML 响应被存成 .xlsx）直接跳过
        continue
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    cached_workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    cached_errors = []
    for sheet in cached_workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    cached_errors.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
    results.append(
        {
            "name": path.name,
            "size": path.stat().st_size,
            "zip_ok": zipfile.ZipFile(path).testzip() is None,
            "sheets": workbook.sheetnames,
            "nonempty": sum(
                1
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            ),
            "formulas": sum(
                1
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ),
            "charts": sum(len(sheet._charts) for sheet in workbook.worksheets),
            "cached_errors": cached_errors,
            "summary_values": {
                "annual_return": cached_workbook["组合分析"]["B22"].value,
                "annual_volatility": cached_workbook["组合分析"]["B23"].value,
                "monthly_var": cached_workbook["组合分析"]["B25"].value,
            },
        }
    )
    if "行业风险归因" in workbook.sheetnames:
        sector = cached_workbook["行业风险归因"]
        backtest = cached_workbook["风险回测"]
        portfolio_var = cached_workbook["组合分析"]["B25"].value
        if sector["B42"].value is None:
            # 空白底稿（任务文件）只检查结构，不检查结果值
            continue
        assert math.isclose(sector["B42"].value, portfolio_var, rel_tol=1e-10)
        assert math.isclose(sector["B43"].value, portfolio_var, rel_tol=1e-10)
        assert math.isclose(sum(sector.cell(30, col).value for col in range(2, 8)), 1.0, rel_tol=1e-10)
        assert backtest["B17"].value == sum(backtest[f"H{row}"].value for row in range(4, 15))
        assert math.isclose(backtest["B18"].value, backtest["B17"].value / 11, rel_tol=1e-10)
        assert all(backtest[f"G{row}"].value >= backtest[f"F{row}"].value for row in range(4, 15))

print(json.dumps(results, ensure_ascii=True, indent=2))
